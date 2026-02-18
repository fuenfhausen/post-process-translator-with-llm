"""Document text extraction and re-injection handlers.

Supports DOCX, PPTX, and XLSX file formats. Each handler can:
  - Extract translatable text segments from a document.
  - Re-inject refined text back into the document preserving formatting.
"""

from __future__ import annotations

import io
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from enum import Enum

import structlog

logger = structlog.get_logger(__name__)


class DocumentFormat(str, Enum):
    """Supported document formats for LLM post-processing."""

    DOCX = "docx"
    PPTX = "pptx"
    XLSX = "xlsx"


# Formats that Document Translation supports but we cannot post-process
TRANSLATE_ONLY_FORMATS = {"pdf", "html", "htm", "txt", "csv", "tsv", "tab", "xlf", "xliff", "msg", "jpg", "png"}


@dataclass
class TextSegment:
    """A single translatable text segment with a location reference.

    The *location* is opaque metadata used by the handler to put the
    refined text back in the right place.
    """

    text: str
    location: dict[str, object] = field(default_factory=dict)


class DocumentHandler(ABC):
    """Base class for document text extraction / re-injection."""

    @abstractmethod
    def extract_segments(self, data: bytes) -> list[TextSegment]:
        """Extract all translatable text segments from *data*."""

    @abstractmethod
    def inject_segments(self, data: bytes, segments: list[TextSegment]) -> bytes:
        """Replace text in *data* with the refined segments, return new bytes."""


# ── DOCX ────────────────────────────────────────────────────────────


class DocxHandler(DocumentHandler):
    """Extract / inject text from Microsoft Word (.docx) documents."""

    def extract_segments(self, data: bytes) -> list[TextSegment]:
        from docx import Document

        doc = Document(io.BytesIO(data))
        segments: list[TextSegment] = []

        for para_idx, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if text:
                segments.append(TextSegment(text=text, location={"type": "paragraph", "index": para_idx}))

        # Also extract from tables
        for tbl_idx, table in enumerate(doc.tables):
            for row_idx, row in enumerate(table.rows):
                for cell_idx, cell in enumerate(row.cells):
                    text = cell.text.strip()
                    if text:
                        segments.append(
                            TextSegment(
                                text=text,
                                location={"type": "table_cell", "table": tbl_idx, "row": row_idx, "cell": cell_idx},
                            )
                        )

        logger.info("docx.extract", segment_count=len(segments))
        return segments

    def inject_segments(self, data: bytes, segments: list[TextSegment]) -> bytes:
        from docx import Document

        doc = Document(io.BytesIO(data))

        # Build lookup maps
        para_map: dict[int, str] = {}
        cell_map: dict[tuple[int, int, int], str] = {}
        for seg in segments:
            loc = seg.location
            if loc.get("type") == "paragraph":
                para_map[int(loc["index"])] = seg.text  # type: ignore[arg-type]
            elif loc.get("type") == "table_cell":
                key = (int(loc["table"]), int(loc["row"]), int(loc["cell"]))  # type: ignore[arg-type]
                cell_map[key] = seg.text

        # Replace paragraph text (preserve runs/formatting of first run)
        for idx, para in enumerate(doc.paragraphs):
            if idx in para_map:
                self._replace_paragraph_text(para, para_map[idx])

        # Replace table cell text
        for (tbl_idx, row_idx, cell_idx), text in cell_map.items():
            try:
                cell = doc.tables[tbl_idx].rows[row_idx].cells[cell_idx]
                if cell.paragraphs:
                    self._replace_paragraph_text(cell.paragraphs[0], text)
            except IndexError:
                logger.warning("docx.inject.cell_miss", table=tbl_idx, row=row_idx, cell=cell_idx)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    @staticmethod
    def _replace_paragraph_text(paragraph: object, new_text: str) -> None:
        """Replace all text in a paragraph while keeping the first run's formatting."""
        from docx.text.paragraph import Paragraph

        para: Paragraph = paragraph  # type: ignore[assignment]
        if not para.runs:
            para.text = new_text
            return
        # Keep first run's style, clear the rest
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ""


# ── PPTX ────────────────────────────────────────────────────────────


class PptxHandler(DocumentHandler):
    """Extract / inject text from Microsoft PowerPoint (.pptx) presentations."""

    def extract_segments(self, data: bytes) -> list[TextSegment]:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))
        segments: list[TextSegment] = []

        for slide_idx, slide in enumerate(prs.slides):
            for shape_idx, shape in enumerate(slide.shapes):
                if not shape.has_text_frame:
                    continue
                for para_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    text = paragraph.text.strip()
                    if text:
                        segments.append(
                            TextSegment(
                                text=text,
                                location={
                                    "type": "pptx_para",
                                    "slide": slide_idx,
                                    "shape": shape_idx,
                                    "para": para_idx,
                                },
                            )
                        )

        logger.info("pptx.extract", segment_count=len(segments))
        return segments

    def inject_segments(self, data: bytes, segments: list[TextSegment]) -> bytes:
        from pptx import Presentation

        prs = Presentation(io.BytesIO(data))

        # Build lookup: (slide, shape, para) -> text
        lookup: dict[tuple[int, int, int], str] = {}
        for seg in segments:
            loc = seg.location
            key = (int(loc["slide"]), int(loc["shape"]), int(loc["para"]))  # type: ignore[arg-type]
            lookup[key] = seg.text

        for slide_idx, slide in enumerate(prs.slides):
            for shape_idx, shape in enumerate(slide.shapes):
                if not shape.has_text_frame:
                    continue
                for para_idx, paragraph in enumerate(shape.text_frame.paragraphs):
                    key = (slide_idx, shape_idx, para_idx)
                    if key in lookup:
                        if paragraph.runs:
                            paragraph.runs[0].text = lookup[key]
                            for run in paragraph.runs[1:]:
                                run.text = ""
                        else:
                            paragraph.text = lookup[key]

        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue()


# ── XLSX ────────────────────────────────────────────────────────────


class XlsxHandler(DocumentHandler):
    """Extract / inject text from Microsoft Excel (.xlsx) workbooks."""

    def extract_segments(self, data: bytes) -> list[TextSegment]:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        segments: list[TextSegment] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row_idx, row in enumerate(ws.iter_rows(), start=1):
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.strip():
                        segments.append(
                            TextSegment(
                                text=cell.value.strip(),
                                location={
                                    "type": "xlsx_cell",
                                    "sheet": sheet_name,
                                    "coordinate": cell.coordinate,
                                },
                            )
                        )

        logger.info("xlsx.extract", segment_count=len(segments))
        return segments

    def inject_segments(self, data: bytes, segments: list[TextSegment]) -> bytes:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))

        for seg in segments:
            loc = seg.location
            sheet_name = str(loc["sheet"])
            coordinate = str(loc["coordinate"])
            try:
                wb[sheet_name][coordinate] = seg.text
            except (KeyError, IndexError):
                logger.warning("xlsx.inject.miss", sheet=sheet_name, coord=coordinate)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()


# ── Factory ─────────────────────────────────────────────────────────


def get_handler(fmt: DocumentFormat) -> DocumentHandler:
    """Return the appropriate handler for the given document format."""
    handlers: dict[DocumentFormat, type[DocumentHandler]] = {
        DocumentFormat.DOCX: DocxHandler,
        DocumentFormat.PPTX: PptxHandler,
        DocumentFormat.XLSX: XlsxHandler,
    }
    return handlers[fmt]()


def detect_format(filename: str) -> DocumentFormat | None:
    """Detect document format from filename extension. Returns None if unsupported for post-processing."""
    ext = filename.rsplit(".", maxsplit=1)[-1].lower() if "." in filename else ""
    try:
        return DocumentFormat(ext)
    except ValueError:
        return None
